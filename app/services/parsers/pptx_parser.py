import os
from pptx import Presentation
from typing import List, Dict, Any, Tuple
from loguru import logger

from app.models.document_element import DocumentElement
from app.services.parsers.base import BaseParser, BaseValidator

class PPTXParser(BaseParser):
    """
    Parser for PowerPoint presentations using python-pptx.
    Preserves slide boundaries.
    """
    def parse(self, file_path: str, **kwargs) -> List[DocumentElement]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"PPTX file not found: {file_path}")
            
        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise ValueError(f"Failed to parse PPTX: {e}")
            
        elements = []
        
        for slide_idx, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        elements.append(DocumentElement(
                            element_type="text",
                            content=text,
                            metadata={
                                "source_type": "pptx",
                                "slide_number": slide_idx + 1,
                            }
                        ))
                elif shape.has_table:
                    for row_idx, row in enumerate(shape.table.rows):
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        if row_text.strip():
                            elements.append(DocumentElement(
                                element_type="table_row",
                                content=row_text,
                                metadata={
                                    "source_type": "pptx",
                                    "slide_number": slide_idx + 1,
                                    "table_row_index": row_idx
                                }
                            ))
                            
        logger.info(f"Successfully extracted {len(elements)} elements from PPTX.")
        return elements

    def chunk(self, elements: List[DocumentElement]) -> List[Dict[str, Any]]:
        """
        Group content by slide.
        """
        chunks = []
        current_chunk_text = ""
        current_slide = None
        chunk_index = 0
        
        for el in elements:
            slide_number = el.metadata.get("slide_number")
            
            if current_slide is not None and slide_number != current_slide:
                chunks.append({
                    "chunk_index": chunk_index,
                    "text_content": current_chunk_text.strip(),
                    "metadata_json": {"source_type": "pptx", "slide_number": current_slide},
                    "page_number": current_slide
                })
                chunk_index += 1
                current_chunk_text = ""
                
            current_slide = slide_number
            current_chunk_text += f"{el.content}\n\n"
            
        if current_chunk_text.strip():
            chunks.append({
                "chunk_index": chunk_index,
                "text_content": current_chunk_text.strip(),
                "metadata_json": {"source_type": "pptx", "slide_number": current_slide},
                "page_number": current_slide
            })
            
        logger.info(f"Successfully grouped PPTX into {len(chunks)} chunks.")
        return chunks

    def get_validator(self) -> BaseValidator:
        return PPTXValidator()


class PPTXValidator(BaseValidator):
    def validate(self, chunks: List[Dict[str, Any]], **kwargs) -> Tuple[str, str, List[Any]]:
        if not chunks:
            return "WARNING: PPTX file is empty or no chunks produced.", "WARNING", []
            
        slides = set(c["metadata_json"].get("slide_number") for c in chunks)
        total_chars = sum(len(c["text_content"]) for c in chunks)
        
        report = []
        report.append("-" * 40)
        report.append("PPTX INGESTION VALIDATION REPORT")
        report.append("-" * 40)
        report.append(f"Total chunks: {len(chunks)}")
        report.append(f"Slides with content: {len(slides)}")
        report.append(f"Total characters: {total_chars}")
        
        status = "PASS" if total_chars > 0 else "WARNING"
            
        report.append(f"\nFINAL RESULT\n------------\n{status}")
        return "\n".join(report), status, []
