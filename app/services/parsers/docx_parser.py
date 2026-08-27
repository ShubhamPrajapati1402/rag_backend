import os
from docx import Document as DocxDocument
from typing import List, Dict, Any, Tuple
from loguru import logger

from app.models.document_element import DocumentElement
from app.services.parsers.base import BaseParser, BaseValidator

class DOCXParser(BaseParser):
    """
    Parser for Word documents using python-docx.
    Preserves heading hierarchy.
    """
    def parse(self, file_path: str, **kwargs) -> List[DocumentElement]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")
            
        try:
            doc = DocxDocument(file_path)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {e}")
            
        elements = []
        heading_path = []
        current_heading_level = 0
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
                
            style_name = para.style.name if para.style else ""
            
            if style_name.startswith('Heading'):
                try:
                    level = int(style_name.split()[-1])
                    current_heading_level = level
                    
                    heading_path = heading_path[:current_heading_level - 1]
                    heading_path.append(text)
                    
                    elements.append(DocumentElement(
                        element_type="title",
                        content=text,
                        metadata={"source_type": "docx", "heading_path": list(heading_path), "level": current_heading_level}
                    ))
                except ValueError:
                    elements.append(DocumentElement(
                        element_type="text",
                        content=text,
                        metadata={"source_type": "docx", "heading_path": list(heading_path)}
                    ))
            else:
                elements.append(DocumentElement(
                    element_type="text",
                    content=text,
                    metadata={"source_type": "docx", "heading_path": list(heading_path)}
                ))
                
        # Also extract tables
        for table in doc.tables:
            for i, row in enumerate(table.rows):
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    elements.append(DocumentElement(
                        element_type="table_row",
                        content=row_text,
                        metadata={"source_type": "docx", "table_row_index": i}
                    ))
                    
        logger.info(f"Successfully extracted {len(elements)} elements from DOCX.")
        return elements

    def chunk(self, elements: List[DocumentElement]) -> List[Dict[str, Any]]:
        """
        Heading-aware chunking for DOCX.
        """
        chunks = []
        current_chunk_text = ""
        current_metadata = {"source_type": "docx"}
        chunk_index = 0
        
        for el in elements:
            meta = el.metadata
            path_str = " > ".join(meta.get("heading_path", [])) if "heading_path" in meta else "Table"
            
            if path_str != current_metadata.get("_path_str") and current_chunk_text.strip():
                chunks.append({
                    "chunk_index": chunk_index,
                    "text_content": current_chunk_text.strip(),
                    "metadata_json": current_metadata,
                    "page_number": None
                })
                chunk_index += 1
                current_chunk_text = ""
            
            if el.element_type == "title":
                current_chunk_text += f"\n\n{el.content}\n"
            else:
                current_chunk_text += f"\n{el.content}"
                
            current_metadata = dict(meta)
            current_metadata["_path_str"] = path_str
            
        if current_chunk_text.strip():
            if "_path_str" in current_metadata:
                del current_metadata["_path_str"]
            chunks.append({
                "chunk_index": chunk_index,
                "text_content": current_chunk_text.strip(),
                "metadata_json": current_metadata,
                "page_number": None
            })
            
        logger.info(f"Successfully grouped DOCX into {len(chunks)} chunks.")
        return chunks

    def get_validator(self) -> BaseValidator:
        return DOCXValidator()


class DOCXValidator(BaseValidator):
    def validate(self, chunks: List[Dict[str, Any]], **kwargs) -> Tuple[str, str, List[Any]]:
        if not chunks:
            return "WARNING: DOCX file is empty or no chunks produced.", "WARNING", []
            
        total_chars = sum(len(c["text_content"]) for c in chunks)
        
        report = []
        report.append("-" * 40)
        report.append("DOCX INGESTION VALIDATION REPORT")
        report.append("-" * 40)
        report.append(f"Total chunks: {len(chunks)}")
        report.append(f"Total characters: {total_chars}")
        
        status = "PASS" if total_chars > 0 else "WARNING"
            
        report.append(f"\nFINAL RESULT\n------------\n{status}")
        return "\n".join(report), status, []
