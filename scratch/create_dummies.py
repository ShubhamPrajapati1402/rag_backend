import os
import json
import pandas as pd
from docx import Document
from pptx import Presentation
import xml.etree.ElementTree as ET

os.makedirs('data/uploads', exist_ok=True)

# Markdown
with open('data/uploads/dummy.md', 'w') as f:
    f.write("# Title\n\nSome paragraph.\n\n## Subtitle\n\nAnother paragraph.")

# TXT
with open('data/uploads/dummy.txt', 'w') as f:
    f.write("Line 1\nLine 2\n\nParagraph 2")

# CSV
pd.DataFrame({"A": [1, 2], "B": [3, 4]}).to_csv('data/uploads/dummy.csv', index=False)

# Excel
with pd.ExcelWriter('data/uploads/dummy.xlsx') as writer:
    pd.DataFrame({"A": [1, 2], "B": [3, 4]}).to_excel(writer, sheet_name='Sheet1', index=False)

# JSON
with open('data/uploads/dummy.json', 'w') as f:
    json.dump({"key": "value", "nested": {"list": [1, 2]}}, f)

# HTML
with open('data/uploads/dummy.html', 'w') as f:
    f.write("<html><body><h1>Title</h1><p>Text</p></body></html>")

# XML
root = ET.Element("root")
child = ET.SubElement(root, "child")
child.text = "value"
tree = ET.ElementTree(root)
tree.write('data/uploads/dummy.xml')

# DOCX
doc = Document()
doc.add_heading('Title', 0)
doc.add_paragraph('Some text.')
doc.save('data/uploads/dummy.docx')

# PPTX
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Title"
prs.save('data/uploads/dummy.pptx')

print("Created dummy files in data/uploads/")
